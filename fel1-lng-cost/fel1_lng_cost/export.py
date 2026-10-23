"""Excel workbook with every Step 13 output, including native Excel charts."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import factors as F
from .estimate import EstimateResult

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)

DISCLAIMER = (
    "FEL 1 (AACE Class 5) screening estimate from public analog projects, following the "
    "FEL 1 LNG Capital Cost Estimating Guideline v0.1. The range is a set of reasoned bounds, "
    "not a statistical P10 and P90. Not for budget authorisation or an investment decision."
)

JOURNEY_STEPS = [
    ("u_reported", "Reported"),
    ("u_escalated", "Step 4 escalated"),
    ("u_located", "Step 5 located"),
    ("u_scoped", "Step 6 scope"),
    ("u_final", "Step 7 size"),
]


def _table(ws, df: pd.DataFrame, start_row: int = 1, widths: dict | None = None) -> int:
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=start_row, column=j, value=str(col))
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for j, v in enumerate(row, start=1):
            if hasattr(v, "item"):
                v = v.item()
            if isinstance(v, float):
                v = round(v, 4)
            if isinstance(v, (list, tuple, dict)):
                v = str(v)
            ws.cell(row=i, column=j, value=v)
    for j, col in enumerate(df.columns, start=1):
        w = (widths or {}).get(col, min(max(len(str(col)) + 2, 10), 45))
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=start_row + 1, column=2)
    return start_row + len(df) + 1


def size_curve(r: EstimateResult, points: int = 21) -> pd.DataFrame:
    """The scaling curve through the combined level at the target train size.

    Values are one-train, Step 6 basis (before the multi-train target factor),
    so they compare directly with each analog's Step 6 value.
    """
    first = r.journey["case"].iloc[0]
    j = r.journey[r.journey["case"] == first]
    u0 = float(r.cases.iloc[0]["unit"]) / float(j["f_target_trains"].iloc[0])
    sizes = r.pool.loc[r.pool["id"].isin(j["id"]), "train_size_mtpa"]
    t0 = r.basis.train_size_mtpa
    lo, hi = min(sizes.min(), t0), max(sizes.max(), t0)
    t = [lo + (hi - lo) * k / (points - 1) for k in range(points)]
    return pd.DataFrame({
        "train_size_mtpa": t,
        "curve_usd_per_tpa": [u0 * (x / t0) ** (r.exponent - 1) for x in t],
    })


def build_workbook(r: EstimateResult, comparison: dict | None = None) -> Workbook:
    b = r.basis
    wb = Workbook()

    # ---------------------------------------------------------- Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "FEL 1 LNG capital cost estimate"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = DISCLAIMER
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:D2")
    ws.row_dimensions[2].height = 60
    fx = b.report_fx
    rows = [
        ("Option", b.option_name),
        ("Purpose", F.PURPOSES.get(b.purpose, b.purpose)),
        ("Facility", F.FACILITY_TYPES[b.facility_type]),
        ("Trains x size (MTPA)", f"{b.trains} x {b.train_size_mtpa:g}"),
        ("Total capacity (MTPA)", b.capacity_mtpa),
        ("Site type", b.site_type),
        ("Location", f"{b.location_name}, factor {b.target_location_factor():.2f}"),
        ("Money of year", b.estimate_year),
        ("Currency", b.report_currency),
        (None, None),
        (f"Low ({b.report_currency} million)", round(r.low * fx)),
        (f"Base ({b.report_currency} million)", round(r.central * fx)),
        (f"High ({b.report_currency} million)", round(r.high * fx)),
        ("Base unit cost (USD per tpa)", round(r.unit_cost)),
        ("Range vs base", f"{r.low / r.central - 1:+.0%} / {r.high / r.central - 1:+.0%}"),
        ("Class", "AACE Class 5"),
        ("Accuracy target", f"{b.class_low:+.0%} / {b.class_high:+.0%}"),
        ("Base method", f"{r.base_method}: {r.base_reason}"),
        ("Analogs used", r.stats["count"]),
        ("Size exponent", round(r.exponent, 3)),
        ("Weighted mean / simple mean / median (USD/tpa)",
         f"{r.stats['weighted_mean']:,.0f} / {r.stats['simple_mean']:,.0f} / {r.stats['median']:,.0f}"),
    ]
    row = 4
    for label, value in rows:
        if label:
            ws.cell(row=row, column=1, value=label).font = BOLD
            ws.cell(row=row, column=2, value=value)
        row += 1
    if r.warnings:
        row += 1
        ws.cell(row=row, column=1, value="Warnings").font = BOLD
        for w in r.warnings:
            row += 1
            ws.cell(row=row, column=1, value=w).alignment = Alignment(wrap_text=True)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            ws.row_dimensions[row].height = 30
    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 55

    # ---------------------------------------------------------- Journey + chart
    wj = wb.create_sheet("Analog journey")
    j = r.journey.merge(r.weights[["id", "weight", "weight_share", "dropped"]], on="id", how="left")
    cols = ["case", "project", "u_reported", "f_esc", "u_escalated", "loc_analog", "loc_target", "f_loc",
            "u_located", "f_split", "f_owner", "f_site", "f_multi", "f_gas", "co2_adj", "f_hybrid",
            "u_scoped", "n_adjustments", "size_ratio", "f_size", "f_target_trains", "u_final",
            "weight", "weight_share", "dropped"]
    end = _table(wj, j[cols], widths={"project": 36})
    first = r.journey["case"].iloc[0]
    chart_rows = r.journey[r.journey["case"] == first]
    wc = wb.create_sheet("Journey chart data")
    data = pd.DataFrame({"step": [s for _, s in JOURNEY_STEPS]})
    for _, a in chart_rows.iterrows():
        data[a["project"][:30]] = [a[k] for k, _ in JOURNEY_STEPS]
    _table(wc, data)
    ch = LineChart()
    ch.title = f"Analog journey, USD per tpa ({first} escalation case)"
    ch.y_axis.title = "USD per tpa"
    ch.height, ch.width = 12, 26
    ch.add_data(Reference(wc, min_col=2, max_col=data.shape[1], min_row=1, max_row=len(data) + 1),
                titles_from_data=True)
    ch.set_categories(Reference(wc, min_col=1, min_row=2, max_row=len(data) + 1))
    wj.add_chart(ch, f"B{end + 3}")

    # ---------------------------------------------------------- Weights
    wwt = wb.create_sheet("Weights")
    wcols = ["project", "same_region", "clean_figure", "size_close", "few_adjustments", "recent",
             "rubric_points", "suggested_weight", "weight", "weight_share", "reason", "value",
             "vs_median", "outlier", "dropped", "drop_reason"]
    end = _table(wwt, r.weights[wcols], widths={"project": 36, "reason": 40})
    wwt.cell(row=end + 2, column=1, value="Weighted mean (USD/tpa)").font = BOLD
    wwt.cell(row=end + 2, column=2, value=round(r.stats["weighted_mean"]))
    wwt.cell(row=end + 3, column=1, value="Simple mean").font = BOLD
    wwt.cell(row=end + 3, column=2, value=round(r.stats["simple_mean"]))
    wwt.cell(row=end + 4, column=1, value="Median").font = BOLD
    wwt.cell(row=end + 4, column=2, value=round(r.stats["median"]))

    # ---------------------------------------------------------- Escalation
    we = wb.create_sheet("Escalation")
    end = _table(we, r.cases)
    if len(r.cases) > 1:
        pivot = r.journey.pivot_table(index="project", columns="case", values="u_final").reset_index()
        start = end + 2
        end2 = _table(we, pivot, start_row=start, widths={"project": 36})
        bc = BarChart()
        bc.title = "Escalation cases per analog (USD per tpa)"
        bc.height, bc.width = 10, 24
        bc.add_data(Reference(we, min_col=2, max_col=pivot.shape[1], min_row=start, max_row=end2),
                    titles_from_data=True)
        bc.set_categories(Reference(we, min_col=1, min_row=start + 1, max_row=end2))
        we.add_chart(bc, f"E{start}")
    if len(r.calibration):
        start = we.max_row + 3
        we.cell(row=start - 1, column=1, value="Damping calibration pairs").font = BOLD
        _table(we, r.calibration, start_row=start, widths={"earlier": 36, "later": 36})

    # ---------------------------------------------------------- Size curve
    wsz = wb.create_sheet("Size curve")
    sc = r.journey[r.journey["case"] == first].merge(
        r.pool[["id", "train_size_mtpa"]], on="id")[["project", "train_size_mtpa", "u_scoped", "u_final"]]
    end = _table(wsz, sc, widths={"project": 36})
    curve = size_curve(r)
    start = end + 2
    end2 = _table(wsz, curve, start_row=start)
    sch = ScatterChart()
    sch.title = f"Analogs (Step 6 value) against train size; curve n = {r.exponent:.2f}"
    sch.x_axis.title = "Train size (MTPA)"
    sch.y_axis.title = "USD per tpa"
    sch.height, sch.width = 12, 22
    pts = Series(Reference(wsz, min_col=3, min_row=2, max_row=end),
                 Reference(wsz, min_col=2, min_row=2, max_row=end), title="Analogs")
    pts.marker.symbol = "circle"
    pts.graphicalProperties.line.noFill = True
    sch.series.append(pts)
    sch.series.append(Series(Reference(wsz, min_col=2, min_row=start + 1, max_row=end2),
                             Reference(wsz, min_col=1, min_row=start + 1, max_row=end2), title="Curve"))
    wsz.add_chart(sch, "G2")

    # ---------------------------------------------------------- Cost blocks
    wb_ = wb.create_sheet("Cost blocks")
    end = _table(wb_, r.blocks_by_area, widths={"block": 34, "note": 30})
    _table(wb_, r.blocks_by_type, start_row=end + 2, widths={"cost_type": 34})

    # ---------------------------------------------------------- Risk register
    wr = wb.create_sheet("Risk register")
    _table(wr, r.site_register, widths={"item": 30, "treatment": 50, "driver": 30, "source": 30})

    # ---------------------------------------------------------- Range
    wg = wb.create_sheet("Range")
    rp = pd.DataFrame([{"component": k, "value": v} for k, v in r.range_parts.items()])
    end = _table(wg, rp, widths={"component": 30})
    rng = pd.DataFrame([{"option": b.option_name, "low": r.low, "base": r.central, "high": r.high}])
    if comparison:
        rng = pd.concat([rng, pd.DataFrame([{
            "option": f"Difference {comparison['option_b']} - {comparison['option_a']}",
            "low": comparison["low"], "base": comparison["central_difference"], "high": comparison["high"],
        }])], ignore_index=True)
    start = end + 2
    end2 = _table(wg, rng, start_row=start, widths={"option": 36})
    rc = BarChart()
    rc.type = "bar"
    rc.title = "Low, base and high (USD million)"
    rc.height, rc.width = 8, 20
    rc.add_data(Reference(wg, min_col=2, max_col=4, min_row=start, max_row=end2), titles_from_data=True)
    rc.set_categories(Reference(wg, min_col=1, min_row=start + 1, max_row=end2))
    wg.add_chart(rc, f"G{start}")
    start = end2 + 2
    _table(wg, r.sensitivities, start_row=start, widths={"sensitivity": 40})

    # ---------------------------------------------------------- Methods, log, references
    wm = wb.create_sheet("Methods")
    end = _table(wm, r.methods, widths={"name": 30, "note": 60})
    wm.cell(row=end + 2, column=1, value="Choice").font = BOLD
    wm.cell(row=end + 2, column=2, value=f"{r.base_method}: {r.base_reason}")
    wm.cell(row=end + 3, column=1, value="Gap explanation").font = BOLD
    wm.cell(row=end + 3, column=2, value=b.method_gap_note or "(none recorded)")

    wl = wb.create_sheet("Decision log")
    _table(wl, r.log, widths={"switch": 40, "setting": 30, "default": 30, "source_of_default": 40, "reason": 40})

    wref = wb.create_sheet("References")
    _table(wref, r.references, widths={"item": 50, "used_for": 40, "source": 70})

    wp = wb.create_sheet("Analog pool")
    pcols = ["id", "project", "country", "region", "site_type", "trains", "capacity_mtpa", "train_size_mtpa",
             "fid_year", "cost_year", "cost_usd_million", "plant_usd_million", "cost_basis", "scope",
             "contingency_status", "process_family", "epc_contractor", "data_quality", "selected",
             "selection_note", "source_url", "notes"]
    _table(wp, r.pool[pcols], widths={"project": 36, "selection_note": 40, "notes": 80, "source_url": 50})
    return wb


def to_excel_bytes(result: EstimateResult, comparison: dict | None = None) -> bytes:
    buf = BytesIO()
    build_workbook(result, comparison).save(buf)
    return buf.getvalue()


def to_excel_file(result: EstimateResult, path: str | Path, comparison: dict | None = None) -> Path:
    path = Path(path)
    build_workbook(result, comparison).save(path)
    return path
